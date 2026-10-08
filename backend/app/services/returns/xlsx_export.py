"""Server-side .xlsx rendering of GSTR-1 / GSTR-3B payloads (task 8.4).

OSS writer: openpyxl, declared in backend/pyproject.toml. One payload per
return is built by ``app.services.returns.pipeline``; this module renders
that payload into a real .xlsx workbook. Money stays integer paise in the
``*_paise`` columns so Excel and the JSON can be compared cell-for-cell;
a rupee display column is derived at render time only (never stored,
never compared — paise → rupees conversion happens here and nowhere else).

Mypy: openpyxl ships no py.typed; types-openpyxl (dev group) provides the
stubs. Import paths below are the public API surface openpyxl documents.
"""

from __future__ import annotations

from decimal import ROUND_HALF_UP, Decimal
from io import BytesIO
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.worksheet.worksheet import Worksheet


def _rupees(paise: int) -> Decimal:
    """Paise → rupees, HALF_UP to 2dp — display only, never a boundary value."""
    return (Decimal(paise) / Decimal(100)).quantize(Decimal("0.01"), ROUND_HALF_UP)


def _header(ws: Worksheet, title: str, payload: dict[str, Any]) -> None:
    ws.append([title])
    ws.append(["GSTIN", payload.get("gstin", "")])
    ws.append(["Filing Period", payload.get("fp", "")])
    ws.append([])


def _sheet(wb: Workbook, title: str) -> Worksheet:
    ws: Worksheet = wb.create_sheet(title)
    return ws


def _remove_default_sheet(wb: Workbook) -> None:
    """Drop the default 'Sheet'; keep mypy happy about the union type."""
    default = wb.active
    if default is not None:
        wb.remove(default)


def _b2b_sheet(wb: Workbook, payload: dict[str, Any]) -> None:
    ws = _sheet(wb, "B2B")
    _header(ws, "GSTR-1 B2B Invoices", payload)
    ws.append(
        [
            "Recipient GSTIN",
            "Invoice No",
            "Invoice Date",
            "Invoice Value (paise)",
            "Place of Supply",
            "Invoice Type",
            "Line",
            "HSN/SAC",
            "Rate %",
            "Taxable Value (paise)",
            "IGST (paise)",
            "CGST (paise)",
            "SGST (paise)",
            "Cess (paise)",
            "Taxable Value (INR)",
        ]
    )
    for entry in payload.get("b2b", []):
        ctin = entry.get("ctin", "")
        for inv in entry.get("inv", []):
            base = [
                ctin,
                inv.get("inum", ""),
                inv.get("idt", ""),
                int(inv.get("val_paise") or 0),
                inv.get("pos", ""),
                inv.get("inv_typ", ""),
            ]
            if not inv.get("itms"):
                ws.append(base + ["-", "", 0, 0, 0, 0, 0, 0])
                continue
            for item in inv.get("itms", []):
                ws.append(
                    base
                    + [
                        int(item.get("num") or 0),
                        item.get("hsn_sac") or "",
                        float(item.get("rt") or 0.0),
                        int(item.get("txval_paise") or 0),
                        int(item.get("iamt_paise") or 0),
                        int(item.get("camt_paise") or 0),
                        int(item.get("samt_paise") or 0),
                        int(item.get("csamt_paise") or 0),
                        _rupees(int(item.get("txval_paise") or 0)),
                    ]
                )


def _cdnr_sheet(wb: Workbook, payload: dict[str, Any]) -> None:
    ws = _sheet(wb, "CDNR")
    _header(ws, "GSTR-1 B2B Credit/Debit Notes", payload)
    ws.append(
        [
            "Recipient GSTIN",
            "Note No",
            "Note Date",
            "Note Type",
            "Reason",
            "Note Value (paise)",
            "HSN/SAC",
            "Rate %",
            "Taxable Value (paise)",
            "IGST (paise)",
            "CGST (paise)",
            "SGST (paise)",
            "Cess (paise)",
        ]
    )
    for entry in payload.get("cdnr", []):
        ctin = entry.get("ctin", "")
        for note in entry.get("nt", []):
            base = [
                ctin,
                note.get("nt_num", ""),
                note.get("nt_dt", ""),
                note.get("nt_typ", ""),
                note.get("rsn", ""),
                int(note.get("val") or 0),
            ]
            if not note.get("itms"):
                ws.append(base + ["", 0, 0, 0, 0, 0, 0])
                continue
            for item in note.get("itms", []):
                ws.append(
                    base
                    + [
                        item.get("hsn_sac") or "",
                        float(item.get("rt") or 0.0),
                        int(item.get("txval_paise") or 0),
                        int(item.get("iamt_paise") or 0),
                        int(item.get("camt_paise") or 0),
                        int(item.get("samt_paise") or 0),
                        int(item.get("csamt_paise") or 0),
                    ]
                )


def _cdnur_sheet(wb: Workbook, payload: dict[str, Any]) -> None:
    ws = _sheet(wb, "CDNUR")
    _header(ws, "GSTR-1 Unregistered Credit/Debit Notes", payload)
    ws.append(
        [
            "Type",
            "Note No",
            "Note Date",
            "Note Type",
            "Reason",
            "Note Value (paise)",
            "Place of Supply",
            "HSN/SAC",
            "Rate %",
            "Taxable Value (paise)",
            "IGST (paise)",
            "CGST (paise)",
            "SGST (paise)",
            "Cess (paise)",
        ]
    )
    for entry in payload.get("cdnur", []):
        typ = entry.get("typ", "")
        for note in entry.get("nt", []):
            base = [
                typ,
                note.get("nt_num", ""),
                note.get("nt_dt", ""),
                note.get("nt_typ", ""),
                note.get("rsn", ""),
                int(note.get("val") or 0),
                note.get("pos", ""),
            ]
            if not note.get("itms"):
                ws.append(base + ["", 0, 0, 0, 0, 0, 0])
                continue
            for item in note.get("itms", []):
                ws.append(
                    base
                    + [
                        item.get("hsn_sac") or "",
                        float(item.get("rt") or 0.0),
                        int(item.get("txval_paise") or 0),
                        int(item.get("iamt_paise") or 0),
                        int(item.get("camt_paise") or 0),
                        int(item.get("samt_paise") or 0),
                        int(item.get("csamt_paise") or 0),
                    ]
                )


def _hsn_sheet(wb: Workbook, rows: list[dict[str, Any]], payload: dict[str, Any]) -> None:
    from app.services.returns.pipeline import hsn_summary_totals

    ws = _sheet(wb, "HSN Summary")
    _header(ws, "GSTR-1 HSN Summary", payload)
    ws.append(
        [
            "HSN/SAC",
            "Rate %",
            "Taxable Value (paise)",
            "IGST (paise)",
            "CGST (paise)",
            "SGST (paise)",
            "Cess (paise)",
            "Num of Lines",
            "Taxable Value (INR)",
        ]
    )
    for row in rows:
        ws.append(
            [
                row["hsn_sac"],
                float(row["rt"]),
                row["txval_paise"],
                row["iamt_paise"],
                row["camt_paise"],
                row["samt_paise"],
                row["csamt_paise"],
                row["num_of_lines"],
                _rupees(row["txval_paise"]),
            ]
        )
    totals = hsn_summary_totals(rows)
    ws.append([])
    ws.append(
        [
            "TOTAL",
            "",
            totals["txval_paise"],
            totals["iamt_paise"],
            totals["camt_paise"],
            totals["samt_paise"],
            totals["csamt_paise"],
            sum(r["num_of_lines"] for r in rows),
            _rupees(totals["txval_paise"]),
        ]
    )


def _doc_issue_sheet(wb: Workbook, payload: dict[str, Any]) -> None:
    ws = _sheet(wb, "Doc Issue")
    _header(ws, "GSTR-1 Document Issue Details", payload)
    ws.append(
        ["Sr", "From", "To", "Total", "Cancelled", "Net Issued"]
    )
    doc_issue = payload.get("doc_issue") or {}
    for rng in doc_issue.get("doc_det", []):
        ws.append(
            [
                int(rng.get("num") or 0),
                rng.get("from_num", ""),
                rng.get("to_num", ""),
                int(rng.get("tot_num") or 0),
                int(rng.get("cancel") or 0),
                int(rng.get("net_issue") or 0),
            ]
        )


def render_gstr1_xlsx(payload: dict[str, Any]) -> bytes:
    """Render the GSTR-1 payload into a real .xlsx workbook."""
    from app.services.returns.pipeline import build_hsn_summary_from_payload

    wb = Workbook()
    _remove_default_sheet(wb)  # drop the default sheet; we create named ones

    _b2b_sheet(wb, payload)
    _cdnr_sheet(wb, payload)
    _cdnur_sheet(wb, payload)
    _hsn_sheet(wb, build_hsn_summary_from_payload(payload), payload)
    _doc_issue_sheet(wb, payload)

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _3b_table_sheet(
    wb: Workbook, title: str, payload: dict[str, Any], table: dict[str, Any]
) -> None:
    ws = _sheet(wb, title)
    _header(ws, f"GSTR-3B {title}", payload)
    ws.append(["Section", "Taxable Value (paise)", "IGST", "CGST", "SGST", "Cess"])
    for section, values in table.items():
        if not isinstance(values, dict):
            continue
        ws.append(
            [
                section,
                int(values.get("txval_paise") or 0),
                int(values.get("iamt_paise") or 0),
                int(values.get("camt_paise") or 0),
                int(values.get("samt_paise") or 0),
                int(values.get("csamt_paise") or 0),
            ]
        )


def render_gstr3b_xlsx(payload: dict[str, Any]) -> bytes:
    """Render the GSTR-3B payload into a real .xlsx workbook."""
    wb = Workbook()
    _remove_default_sheet(wb)

    _3b_table_sheet(wb, "Supplies", payload, payload.get("sup_details", {}))
    _3b_table_sheet(wb, "Inter-State Supplies", payload, payload.get("inter_sup", {}))

    itc_elg = payload.get("itc_elg", {})
    ws = _sheet(wb, "ITC")
    _header(ws, "GSTR-3B ITC Available", payload)
    ws.append(["Type", "IGST (paise)", "CGST (paise)", "SGST (paise)", "Cess (paise)"])
    for item in itc_elg.get("itc_avl", []):
        ws.append(
            [
                item.get("ty", ""),
                int(item.get("iamt_paise") or 0),
                int(item.get("camt_paise") or 0),
                int(item.get("samt_paise") or 0),
                int(item.get("csamt_paise") or 0),
            ]
        )

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def render_return_xlsx(form: str, payload: dict[str, Any]) -> bytes:
    """Dispatch the xlsx renderer by return form."""
    if form == "gstr1":
        return render_gstr1_xlsx(payload)
    if form == "gstr3b":
        return render_gstr3b_xlsx(payload)
    raise ValueError(f"unsupported return form: {form}")


def open_xlsx_bytes(data: bytes) -> Workbook:
    """Test/helper hook: load rendered bytes back as a workbook."""
    return load_workbook(BytesIO(data))
