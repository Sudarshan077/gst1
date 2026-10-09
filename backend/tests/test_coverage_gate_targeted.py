"""Task 10.3 — targeted coverage for the remaining uncovered backend lines.

Why direct calls: on this Windows / CPython-3.11 host the coverage CTracer
under-attributes statements executed downstream of a suspended FastAPI
async-generator dependency (the ``get_session`` yield), so API-driven suites
read service/route bodies as uncovered even though the requests succeed
(see tests/test_review_service.py:1-8 for the same diagnosis). These tests
drive the *handler functions and services directly* — the sanctioned remedy
already used by test_validation_service_branches_direct / test_review_service —
so the coverage gate reflects code the suite genuinely executes.

Nothing here fabricates GSTINs/PANs (every fixture GSTIN comes from
make_gstin, mod-36) and every money value is integer paise.
"""

from __future__ import annotations

import io
import uuid
from datetime import date
from decimal import Decimal
from typing import Any

import pytest
from app.api.errors import ServiceError
from app.core.access import GstinAccess, resolve_gstin_access
from app.core.auth import redis_client as redis_client_mod
from app.core.auth.otp import request_otp
from app.db.models.gst import (
    CreditDebitNote,
    Gstr2bSource,
    Invoice,
    InvoiceDirection,
    InvoiceLine,
    InvoiceStatus,
    NoteStatus,
    NoteType,
    SupplyType,
)
from httpx import AsyncClient
from sqlalchemy.ext.asyncio import async_sessionmaker
from starlette.requests import Request

from tests.gstin_fixtures import gstin_checksum_valid, make_gstin
from tests.v4_helpers import seed_account, seed_open_period

SessionMaker = async_sessionmaker[Any]

FP = "102026"


# --------------------------------------------------------------------- helpers


def _request(headers: dict[str, str]) -> Request:
    scope = {
        "type": "http",
        "method": "POST",
        "path": "/x",
        "headers": [(k.lower().encode(), v.encode()) for k, v in headers.items()],
    }
    return Request(scope)


def _user_id(tokens: dict[str, Any]) -> str:
    from app.core.auth.tokens import verify_access_token

    return str(verify_access_token(tokens["access_token"]))


async def _access_for(
    sessionmaker: SessionMaker, gstin: str, user_id: uuid.UUID
) -> GstinAccess:
    async with sessionmaker() as session:
        return await resolve_gstin_access(session, user_id, gstin)


async def _seed_purchase_invoice(
    sessionmaker: SessionMaker,
    gstin: str,
    fp: str,
    *,
    invoice_no: str,
    supplier_gstin: str,
    total_value_minor: int,
    cgst_minor: int = 0,
    sgst_minor: int = 0,
    igst_minor: int = 0,
) -> uuid.UUID:
    async with sessionmaker() as session:
        inv = Invoice(
            gstin=gstin,
            fp=fp,
            direction=InvoiceDirection.PURCHASE,
            supplier_gstin=supplier_gstin,
            invoice_no=invoice_no,
            invoice_date=date(2026, 10, 1),
            place_of_supply=gstin[:2],
            supply_type=SupplyType.INTRA,
            total_value_minor=total_value_minor,
            status=InvoiceStatus.CONFIRMED,
        )
        session.add(inv)
        await session.flush()
        session.add(
            InvoiceLine(
                invoice_id=inv.id,
                line_no=1,
                gst_rate=Decimal("18.00"),
                taxable_value_minor=total_value_minor - (cgst_minor + sgst_minor),
                cgst_minor=cgst_minor,
                sgst_minor=sgst_minor,
                igst_minor=igst_minor,
                cess_minor=0,
            )
        )
        await session.commit()
        return inv.id


# ------------------------------------------------------- itc router handlers


async def test_itc_router_handlers_direct(
    client: AsyncClient, api_sessionmaker: SessionMaker
) -> None:
    """POST /gstr2b/import, /gstr2b/fetch and GET /itc/report bodies."""
    from app.api.routers import itc as itc_router
    from app.api.schemas_gst import Gstr2bImportIn

    tokens, account = await seed_account(client, api_sessionmaker, state_code="27")
    gstin = account.gstin
    user_id = uuid.UUID(_user_id(tokens))
    supplier = make_gstin(state_code="29")
    await seed_open_period(api_sessionmaker, gstin, FP)
    await _seed_purchase_invoice(
        api_sessionmaker,
        gstin,
        FP,
        invoice_no="ITC-DS-1",
        supplier_gstin=supplier,
        total_value_minor=11800,
        cgst_minor=900,
        sgst_minor=900,
    )

    access = await _access_for(api_sessionmaker, gstin, user_id)

    body = Gstr2bImportIn(
        payload={
            "entries": [
                {
                    "supplier_gstin": supplier,
                    "invoice_no": "ITC-DS-1",
                    "invoice_date": "2026-10-01",
                    "taxable_value_minor": 11800,
                    "cgst_minor": 900,
                    "sgst_minor": 900,
                    "doc_type": "INV",
                }
            ]
        }
    )

    async with api_sessionmaker() as session:
        out = await itc_router.import_2b_statement(gstin, FP, body, access, session, user_id)
    assert out["success"] is True
    assert out["data"].gstin == gstin

    # fetch via GSP returns the stored statement (existing/idempotent branch).
    async with api_sessionmaker() as session:
        fetched = await itc_router.fetch_2b_via_gsp(gstin, FP, access, session, user_id)
    assert fetched["success"] is True
    assert fetched["data"].source == "PORTAL_UPLOAD"

    async with api_sessionmaker() as session:
        recon = await itc_router.reconcile_itc_data(gstin, FP, access, session)
    assert recon["success"] is True

    async with api_sessionmaker() as session:
        report = await itc_router.get_itc_report(gstin, FP, access, session, status=None)
    assert report["success"] is True
    assert report["data"]["summary"]["counts"]["total"] >= 1


# -------------------------------------------------- einvoice router handlers


async def test_einvoice_router_handlers_direct(
    client: AsyncClient, api_sessionmaker: SessionMaker
) -> None:
    """POST /invoices/{id}/irn, /einvoices/{id}/cancel, GET /einvoices bodies."""
    from app.api.routers import einvoice as einv_router

    tokens, account = await seed_account(
        client, api_sessionmaker, state_code="27", aato_minor=7_000_000_000
    )
    gstin = account.gstin
    user_id = uuid.UUID(_user_id(tokens))
    await seed_open_period(api_sessionmaker, gstin, FP)

    async with api_sessionmaker() as session:
        inv = Invoice(
            gstin=gstin,
            fp=FP,
            direction=InvoiceDirection.SALES,
            supplier_gstin=gstin,
            buyer_gstin=make_gstin(state_code="29"),
            invoice_no="EINV-DS-1",
            invoice_date=date(2026, 10, 1),
            place_of_supply="27",
            supply_type=SupplyType.INTRA,
            total_value_minor=118000,
            status=InvoiceStatus.CONFIRMED,
        )
        session.add(inv)
        await session.commit()
        await session.refresh(inv)
        inv_id = inv.id

    async with api_sessionmaker() as session:
        gen = await einv_router.generate_irn(inv_id, session, user_id)
    assert gen["success"] is True
    assert gen["data"]["irn"]

    # missing invoice -> 404 envelope body
    async with api_sessionmaker() as session:
        with pytest.raises(ServiceError) as exc:
            await einv_router.generate_irn(uuid.uuid4(), session, user_id)
    assert exc.value.status_code == 404

    access = await _access_for(api_sessionmaker, gstin, user_id)
    async with api_sessionmaker() as session:
        listed = await einv_router.get_einvoices(gstin, FP, access, session)
    assert listed["success"] is True
    assert len(listed["data"]) == 1

    async with api_sessionmaker() as session:
        cancelled = await einv_router.cancel_irn(inv_id, session, user_id)
    assert cancelled["success"] is True
    assert cancelled["data"]["cancelled_at"] is not None


# ------------------------------------------------ gst_accounts router handlers


async def test_gst_accounts_router_handlers_direct(
    client: AsyncClient, api_sessionmaker: SessionMaker
) -> None:
    from app.api.routers import gst_accounts as ga_router
    from app.api.schemas_gst import GstAccountCreateIn, GstAccountPatchIn

    tokens, account = await seed_account(client, api_sessionmaker, state_code="27")
    gstin = account.gstin
    user_id = uuid.UUID(_user_id(tokens))

    async with api_sessionmaker() as session:
        created = await ga_router.create_gst_account(
            GstAccountCreateIn(gstin=make_gstin(state_code="29"), legal_name="DS Co"),
            session,
            user_id,
        )
    assert created["success"] is True

    async with api_sessionmaker() as session:
        listed = await ga_router.list_gst_accounts(session, user_id)
    assert listed["success"] is True and len(listed["data"]) == 2

    access = await _access_for(api_sessionmaker, gstin, user_id)

    async with api_sessionmaker() as session:
        detail = await ga_router.get_gst_account(access, session)
    assert detail["data"]["gstin"] == gstin

    async with api_sessionmaker() as session:
        patched = await ga_router.patch_gst_account(
            GstAccountPatchIn(trade_name="Renamed Co"), access, session
        )
    assert patched["data"]["trade_name"] == "Renamed Co"

    async with api_sessionmaker() as session:
        overview = await ga_router.get_gst_account_overview(access, session)
    assert overview["data"]["gstin"] == gstin
    assert len(overview["data"]["periods"]) == 12

    await seed_open_period(api_sessionmaker, gstin, FP)
    async with api_sessionmaker() as session:
        summary = await ga_router.get_month_summary(gstin, FP, access, session)
    assert summary["data"]["fp"] == FP

    async with api_sessionmaker() as session:
        audit = await ga_router.get_gstin_audit(access, session)
    assert audit["success"] is True and len(audit["data"]) >= 1

    async with api_sessionmaker() as session:
        collabs = await ga_router.list_collaborators(access, session)
    assert collabs["success"] is True and collabs["data"][0]["role"] == "ADMIN"


# ------------------------------------------------------ gstr2b service (full)


async def test_gstr2b_service_full_paths_direct(
    client: AsyncClient, api_sessionmaker: SessionMaker
) -> None:
    """Direct import -> reconcile (every status) -> build_itc_report (incl. filter)."""
    from app.services.returns import gstr2b

    tokens, account = await seed_account(client, api_sessionmaker, state_code="27")
    gstin = account.gstin
    user_id = uuid.UUID(_user_id(tokens))
    supplier = make_gstin(state_code="29")
    await seed_open_period(api_sessionmaker, gstin, FP)

    # Books: exact match, small-delta probable, and a books-only row.
    await _seed_purchase_invoice(
        api_sessionmaker,
        gstin,
        FP,
        invoice_no="RS-MATCH",
        supplier_gstin=supplier,
        total_value_minor=11800,
        cgst_minor=900,
        sgst_minor=900,
    )
    await _seed_purchase_invoice(
        api_sessionmaker,
        gstin,
        FP,
        invoice_no="RS-DELTA",
        supplier_gstin=supplier,
        total_value_minor=20000,
        cgst_minor=1800,
        sgst_minor=1800,
    )
    await _seed_purchase_invoice(
        api_sessionmaker,
        gstin,
        FP,
        invoice_no="RS-BOOKSPNLY",
        supplier_gstin=supplier,
        total_value_minor=50000,
        cgst_minor=4500,
        sgst_minor=4500,
    )

    payload = {
        "entries": [
            {
                "supplier_gstin": supplier,
                "invoice_no": "RS-MATCH",
                "invoice_date": "2026-10-01",
                "taxable_value_minor": 11800,
                "cgst_minor": 900,
                "sgst_minor": 900,
                "doc_type": "INV",
            },
            {
                "supplier_gstin": supplier,
                "invoice_no": "RS-DELTA",
                "invoice_date": "2026-10-02",
                "taxable_value_minor": 20050,
                "cgst_minor": 1800,
                "sgst_minor": 1800,
                "doc_type": "INV",
            },
            {
                "supplier_gstin": supplier,
                "invoice_no": "RS-2BONLY",
                "invoice_date": "2026-10-05",
                "taxable_value_minor": 30000,
                "cgst_minor": 2700,
                "sgst_minor": 2700,
                "doc_type": "INV",
            },
        ]
    }

    async with api_sessionmaker() as session:
        statement = await gstr2b.import_gstr2b(session, gstin, FP, payload, user_id)
        assert statement.source == Gstr2bSource.PORTAL_UPLOAD

        recons = await gstr2b.reconcile_itc(session, gstin, FP)
        statuses = {r.match_status.value for r in recons}
        assert statuses == {
            "MATCHED",
            "PROBABLE",
            "MISSING_IN_2B",
            "MISSING_IN_BOOKS",
        }

        report = await gstr2b.build_itc_report(session, gstin, FP)
        assert report["statement"]["entry_count"] == 3
        assert report["summary"]["counts"]["total"] == len(recons)
        # books side carries integer-paise buckets on both sides
        assert isinstance(report["summary"]["books_itc_paise"]["total_paise"], int)
        assert isinstance(report["summary"]["gstr2b_itc_paise"]["total_paise"], int)

        filtered = await gstr2b.build_itc_report(session, gstin, FP, "MISSING_IN_BOOKS")
        assert [r["match_status"] for r in filtered["rows"]] == ["MISSING_IN_BOOKS"]
        # counts always cover every status even when rows are filtered
        assert filtered["summary"]["counts"]["total"] == len(recons)

        with pytest.raises(ServiceError) as exc:
            await gstr2b.build_itc_report(session, gstin, FP, "NOPE")
        assert exc.value.status_code == 422


async def test_gstr2b_reconcile_without_statement_direct(
    client: AsyncClient, api_sessionmaker: SessionMaker
) -> None:
    """reconcile_itc early-return when no 2B statement exists."""
    from app.services.returns import gstr2b

    tokens, account = await seed_account(client, api_sessionmaker, state_code="27")
    async with api_sessionmaker() as session:
        assert await gstr2b.reconcile_itc(session, account.gstin, FP) == []


# --------------------------------------------------------- xlsx export edges


def test_xlsx_export_cdnur_and_empty_itms_branches() -> None:
    """CDNUR sheet + the no-itms fallback rows in B2B/CDNR/CDNUR."""
    from app.services.returns.xlsx_export import (
        open_xlsx_bytes,
        render_gstr1_xlsx,
        render_gstr3b_xlsx,
        render_return_xlsx,
    )
    from openpyxl import load_workbook

    gstin = make_gstin()
    payload: dict[str, Any] = {
        "gstin": gstin,
        "fp": FP,
        "b2b": [
            {
                "ctin": gstin,
                "inv": [
                    {
                        "inum": "X1",
                        "idt": "01-10-2026",
                        "val_paise": 100,
                        "pos": "27",
                        "inv_typ": "R",
                        "itms": [],
                    }
                ],
            }
        ],
        "cdnr": [
            {
                "ctin": gstin,
                "nt": [
                    {
                        "nt_num": "C1",
                        "nt_dt": "01-10-2026",
                        "nt_typ": "C",
                        "rsn": "R1",
                        "val": 50,
                        "itms": [],
                    }
                ],
            }
        ],
        "cdnur": [
            {
                "typ": "B2CL",
                "nt": [
                    {
                        "nt_num": "U1",
                        "nt_dt": "01-10-2026",
                        "nt_typ": "C",
                        "rsn": "R1",
                        "val": 10,
                        "pos": "07",
                        "itms": [],
                    },
                    {
                        "nt_num": "U2",
                        "nt_dt": "02-10-2026",
                        "nt_typ": "D",
                        "rsn": "R2",
                        "val": 20,
                        "pos": "07",
                        "itms": [
                            {
                                "hsn_sac": "9983",
                                "rt": 18.0,
                                "txval_paise": 20,
                                "iamt_paise": 0,
                                "camt_paise": 1,
                                "samt_paise": 1,
                                "csamt_paise": 0,
                            }
                        ],
                    },
                ],
            }
        ],
        "doc_issue": {
            "doc_det": [
                {
                    "num": 1,
                    "from_num": "a",
                    "to_num": "b",
                    "tot_num": 1,
                    "cancel": 0,
                    "net_issue": 1,
                }
            ]
        },
    }
    data = render_gstr1_xlsx(payload)
    wb = open_xlsx_bytes(data)
    assert "CDNUR" in wb.sheetnames
    cdnur = list(wb["CDNUR"].iter_rows(min_row=6, values_only=True))
    assert any(r[1] == "U1" for r in cdnur)
    assert any(r[1] == "U2" and r[7] == "9983" for r in cdnur)

    # GSTR-3B with a non-dict section value exercises the isinstance guard.
    g3b = {
        "gstin": gstin,
        "fp": FP,
        "sup_details": {"osup_det": {"txval_paise": 5}},
        "inter_sup": {"bad": "not-a-dict"},
    }
    wb3b = load_workbook(io.BytesIO(render_gstr3b_xlsx(g3b)))
    assert "Supplies" in wb3b.sheetnames

    with pytest.raises(ValueError, match="unsupported return form"):
        render_return_xlsx("nope", payload)


# ---------------------------------------------------- validation service edges


async def test_validation_note_party_gstin_and_zero_value_direct(
    client: AsyncClient, api_sessionmaker: SessionMaker
) -> None:
    """Blocking note-party GSTIN + zero-value warning via direct service call."""
    from app.services.returns.validation import validate_period_for_filing

    tokens, account = await seed_account(client, api_sessionmaker, state_code="27")
    gstin = account.gstin
    await seed_open_period(api_sessionmaker, gstin, FP)

    bad_buyer = make_gstin(state_code="29")
    bad_buyer = bad_buyer[:14] + ("0" if bad_buyer[14] != "0" else "1")
    assert not gstin_checksum_valid(bad_buyer)

    async with api_sessionmaker() as session:
        session.add(
            CreditDebitNote(
                gstin=gstin,
                fp=FP,
                note_type=NoteType.CDN,
                reason_code="01",
                buyer_gstin=bad_buyer,
                taxable_value_minor=0,
                note_no="VAL-NOTE-1",
                note_date=date(2026, 10, 1),
                status=NoteStatus.CONFIRMED,
            )
        )
        await session.commit()

    async with api_sessionmaker() as session:
        result = await validate_period_for_filing(session, gstin, FP)

    rules = [b["rule"] for b in result["blocking"]]
    assert "INVALID_PARTY_GSTIN" in rules
    warn_rules = [w["rule"] for w in result["warnings"]]
    assert "ZERO_VALUE_LINES" in warn_rules


# ------------------------------------------------------- deadline service edges


async def test_deadline_scheme_branches_and_missing_account_direct(
    client: AsyncClient, api_sessionmaker: SessionMaker
) -> None:
    from app.services.returns import deadline

    gstin = make_gstin(state_code="27")
    comp = deadline.compute_due_dates(gstin, "032026", "COMPOSITION")
    assert comp["gstr1_due_date"] is None
    assert comp["gstr4_due_date"] == date(2026, 4, 30)  # Mar quarter -> 30 Apr

    qrmp = deadline.compute_due_dates(gstin, "032026", "QRMP")
    assert qrmp["iff_eligible"] is False  # quarter-end month

    regular = deadline.compute_due_dates(gstin, "122026", "REGULAR_MONTHLY")
    assert regular["gstr1_due_date"] == date(2027, 1, 11)  # year rollover

    for bad_fp in ("13", "132026"):
        with pytest.raises(ValueError):
            deadline.compute_due_dates(gstin, bad_fp, "REGULAR_MONTHLY")

    missing = make_gstin()
    async with api_sessionmaker() as session:
        with pytest.raises(ServiceError) as exc:
            await deadline.list_periods_for_fy(session, missing)
    assert exc.value.status_code == 404

    tokens, account = await seed_account(client, api_sessionmaker, state_code="27")
    async with api_sessionmaker() as session:
        # fy omitted -> derives the current FY and creates its 12 periods
        periods = await deadline.list_periods_for_fy(session, account.gstin)
    assert len(periods) == 12

    async with api_sessionmaker() as session:
        with pytest.raises(ServiceError):
            await deadline.get_month_summary(session, missing, FP)


# ------------------------------------------------------ gstr1 / gstr3b edges


def test_gstr1_notes_cdnr_and_cdnur_and_validator() -> None:
    from app.services.returns.gstr1 import (
        build_gstr1_from_invoices,
        generate_nil_gstr1,
        validate_gstr1_payload,
    )

    gstin = make_gstin(state_code="27")
    buyer = make_gstin(state_code="29")

    class _Note:
        def __init__(self, buyer_gstin: str | None, taxable: int) -> None:
            self.buyer_gstin = buyer_gstin
            self.note_no = "N-1"
            self.note_date = date(2026, 10, 1)
            self.note_type = "CDN"
            self.reason_code = "01"
            self.taxable_value_minor = taxable
            self.cgst_minor = 900
            self.sgst_minor = 900
            self.igst_minor = 0
            self.cess_minor = 0

    # One CDNR note (buyer present) and one zero-value CDNUR note (no buyer).
    payload = build_gstr1_from_invoices(
        gstin, FP, [], notes=[_Note(buyer, 10000), _Note(None, 0)]
    )
    assert payload["cdnr"][0]["ctin"] == buyer
    assert payload["cdnur"][0]["typ"] == "B2CL"
    assert payload["cdnur"][0]["nt"][0]["pos"] == "07"

    # Self-validator failure branch.
    ok, err = validate_gstr1_payload({"gstin": "BAD", "fp": "1"})
    assert ok is False and err

    ok_nil, err_nil = validate_gstr1_payload(generate_nil_gstr1(gstin, FP))
    assert ok_nil is True and err_nil is None


def test_gstr3b_inter_state_unregistered_and_validator() -> None:
    from app.services.returns.gstr3b import (
        build_gstr3b_from_invoices,
        validate_gstr3b_payload,
    )

    gstin = make_gstin(state_code="27")

    class _Line:
        igst_minor = 1800
        cgst_minor = 0
        sgst_minor = 0
        cess_minor = 0

    class _Inv:
        direction = InvoiceDirection.SALES
        supply_type = SupplyType.INTER
        buyer_gstin = None  # unregistered inter-state -> table 3.2 unreg bucket
        total_value_minor = 11800
        rchrg = False
        lines = [_Line()]

    payload = build_gstr3b_from_invoices(gstin, FP, [_Inv()])
    assert payload["inter_sup"]["unreg_det"]["txval_paise"] == 11800
    assert payload["inter_sup"]["unreg_det"]["iamt_paise"] == 1800

    ok, err = validate_gstr3b_payload({"gstin": "BAD", "fp": "1"})
    assert ok is False and err

    ok_good, err_good = validate_gstr3b_payload(payload)
    assert ok_good is True and err_good is None


# ------------------------------------------- auth dependencies / error handlers


async def test_stepup_otp_paths_direct(
    fake_redis: Any, api_sessionmaker: SessionMaker
) -> None:
    from app.core.auth.dependencies import optional_user, require_stepup
    from app.core.auth.errors import StepUpRequired, TokenInvalid
    from app.core.auth.tokens import create_stepup_token
    from app.db.models.core import User

    assert await optional_user(None) is None

    bad = type("C", (), {"credentials": "not-a-jwt", "scheme": "Bearer"})()
    assert await optional_user(bad) is None

    async with api_sessionmaker() as session:
        user = User(email=f"stepup+{uuid.uuid4().hex[:8]}@example.com", full_name="S")
        session.add(user)
        await session.commit()
        await session.refresh(user)
        user_id = user.id
        email = user.email
        assert email is not None

    otp_req = await request_otp(redis_client_mod.get_redis(), email, "LOGIN")
    assert otp_req.dev_otp

    async with api_sessionmaker() as session:
        returned = await require_stepup(
            _request({"X-OTP": otp_req.dev_otp or ""}), session, user_id
        )
    assert returned == user_id

    # no step-up proof at all -> StepUpRequired
    async with api_sessionmaker() as session:
        with pytest.raises(StepUpRequired):
            await require_stepup(_request({}), session, user_id)

    # mismatched step-up token -> TokenInvalid
    async with api_sessionmaker() as session:
        with pytest.raises(TokenInvalid):
            await require_stepup(
                _request({"X-Stepup-Token": create_stepup_token(uuid.uuid4())}),
                session,
                user_id,
            )


async def test_error_handlers_direct() -> None:
    from app.api.errors import (
        http_exception_handler,
        service_error_handler,
        validation_error_handler,
        value_error_handler,
    )
    from fastapi import HTTPException
    from fastapi.exceptions import RequestValidationError

    req = _request({})

    ve = await value_error_handler(req, ValueError("bad value"))
    assert ve.status_code == 422

    se = await service_error_handler(req, ServiceError("nope", 409, "CONFLICT"))
    assert se.status_code == 409

    he = await http_exception_handler(req, HTTPException(status_code=418, detail="teapot"))
    assert he.status_code == 418

    val = await validation_error_handler(req, RequestValidationError([]))
    assert val.status_code == 422


def test_gstin_checksum_invalid_inputs() -> None:
    from app.core.gstin import gstin_checksum_valid

    assert gstin_checksum_valid("") is False
    assert gstin_checksum_valid("27AAPFU0939F1Z") is False  # short -> format reject
    assert gstin_checksum_valid(make_gstin()) is True
